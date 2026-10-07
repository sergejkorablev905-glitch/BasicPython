from openpyxl import Workbook

wb = Workbook()

ws = wb.active
ws.title = "Аналитика_2026"
ws_archive = wb.create_sheet(title="Архив_Данных")

ws["A1"] = "Итоговая выручка"

wb.save("report.xlsx")